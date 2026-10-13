"""Text extraction for scope documents (.docx, .pdf, .txt, .md).

Every block of text is given a short, stable **segment id** so the extraction step
can cite exactly where each scope item came from, and the review gate can verify
that citation deterministically:

    P12        paragraph 12 of a Word document
    T2.R5      table 2, row 5 of a Word document (cells joined with " | ")
    PG3.4      page 3, block 4 of a PDF
    L7         block 7 of a plain-text file

Design rules
------------
* No OCR. A PDF without a text layer (a scan) is rejected with a clear message,
  because unreadable scope must never be silently "filled in" downstream.
* Nothing is written to disk; the document lives only in memory for the session.
* The reader is deterministic: the same file always gives the same segments.
"""
from __future__ import annotations

import io
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Optional

MAX_CHARS = 600_000          # ~150k tokens: well inside Gemini 2.5 Flash context, keeps cost bounded
MIN_TEXT_CHARS = 40          # fewer characters than this = nothing usable was read
SUPPORTED = (".docx", ".pdf", ".txt", ".md")


class DocumentReadError(ValueError):
    """Raised when a document cannot be turned into usable text."""


@dataclass
class Segment:
    sid: str                 # segment id, e.g. "P12", "T2.R5", "PG3.4"
    text: str
    kind: str = "paragraph"  # paragraph | heading | table_row | page_block
    level: int = 0           # heading level (1 = Heading 1) when kind == "heading"


@dataclass
class DocumentText:
    filename: str
    file_type: str
    segments: list = field(default_factory=list)
    pages: Optional[int] = None
    tables: int = 0
    warnings: list = field(default_factory=list)
    truncated: bool = False

    # ------------------------------------------------------------------ #
    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.segments)

    @property
    def word_count(self) -> int:
        return sum(len(s.text.split()) for s in self.segments)

    def by_id(self) -> dict:
        return {s.sid: s for s in self.segments}

    def as_prompt_text(self) -> str:
        """Text sent to the model: one segment per line, prefixed by its id."""
        lines = []
        for s in self.segments:
            tag = f"[{s.sid}]"
            if s.kind == "heading":
                lines.append(f"{tag} {'#' * max(1, min(s.level, 6))} {s.text}")
            else:
                lines.append(f"{tag} {s.text}")
        return "\n".join(lines)

    def preview(self, limit: int = 40) -> list:
        return [(s.sid, s.text[:160]) for s in self.segments[:limit]]


# --------------------------------------------------------------------------- #
# Text normalisation (shared with citation checking in extractor.py)
# --------------------------------------------------------------------------- #
_WS = re.compile(r"\s+")
_QUOTES = {"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-",
           " ": " ", "•": "-", "…": "..."}


def clean(text: str) -> str:
    """Readable clean-up used for stored segment text."""
    if text is None:
        return ""
    t = unicodedata.normalize("NFKC", str(text))
    for k, v in _QUOTES.items():
        t = t.replace(k, v)
    t = t.replace("\x00", "")
    return _WS.sub(" ", t).strip()


def normalise_for_match(text: str) -> str:
    """Aggressive normalisation for verifying quotes: case, punctuation and spacing are ignored."""
    t = clean(text).lower()
    t = re.sub(r"[^a-z0-9%./]+", " ", t)
    return _WS.sub(" ", t).strip()


# --------------------------------------------------------------------------- #
# Readers
# --------------------------------------------------------------------------- #
def _read_docx(data: bytes, doc: DocumentText):
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - dependency missing
        raise DocumentReadError("python-docx is not installed (see requirements.txt).") from exc
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentReadError(f"Not a readable Word document (.docx): {exc}") from exc

    p_no, t_no = 0, 0
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, d)
            text = clean(para.text)
            if not text:
                continue
            p_no += 1
            style = (para.style.name if para.style is not None else "") or ""
            m = re.match(r"(?i)^(heading|titre|überschrift)\s*(\d)", style)
            if m:
                doc.segments.append(Segment(f"P{p_no}", text, "heading", int(m.group(2))))
            elif style.lower() == "title":
                doc.segments.append(Segment(f"P{p_no}", text, "heading", 1))
            elif _looks_like_list(para) or style.lower().startswith("list"):
                doc.segments.append(Segment(f"P{p_no}", "- " + text, "paragraph"))
            else:
                doc.segments.append(Segment(f"P{p_no}", text, "paragraph"))
        elif tag == "tbl":
            t_no += 1
            table = Table(child, d)
            for r_no, row in enumerate(table.rows, start=1):
                cells, last = [], None
                for cell in row.cells:
                    # merged cells repeat the same object; keep each once
                    if last is not None and cell._tc is last:
                        continue
                    last = cell._tc
                    v = clean(cell.text)
                    cells.append(v)
                if any(cells):
                    doc.segments.append(Segment(f"T{t_no}.R{r_no}", " | ".join(cells), "table_row"))
    doc.tables = t_no
    # headers/footers rarely carry scope, but the document title often sits there
    try:
        sec = d.sections[0]
        head = clean(" ".join(p.text for p in sec.header.paragraphs))
        if head and len(head) < 200:
            doc.segments.insert(0, Segment("H1", head, "paragraph"))
    except Exception:  # noqa: BLE001
        pass


def _looks_like_list(para) -> bool:
    try:
        ppr = para._p.pPr
        return ppr is not None and ppr.numPr is not None
    except Exception:  # noqa: BLE001
        return False


def _read_pdf(data: bytes, doc: DocumentText):
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentReadError("pypdf is not installed (see requirements.txt).") from exc
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentReadError(f"Not a readable PDF: {exc}") from exc
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise DocumentReadError("The PDF is password-protected. Remove the password and upload again.")
        except DocumentReadError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise DocumentReadError("The PDF is password-protected. Remove the password and upload again.") from exc

    doc.pages = len(reader.pages)
    empty_pages = 0
    for pg_no, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            raw = ""
        blocks = _split_blocks(raw)
        if not blocks:
            empty_pages += 1
            continue
        for b_no, block in enumerate(blocks, start=1):
            doc.segments.append(Segment(f"PG{pg_no}.{b_no}", block, "page_block"))
    if doc.pages and empty_pages == doc.pages:
        raise DocumentReadError(
            "No text layer found in this PDF (it looks like a scan). Scanned documents are out of scope for the "
            "MVP: run OCR first or upload the original Word/PDF export.")
    if empty_pages:
        doc.warnings.append(f"{empty_pages} of {doc.pages} PDF pages had no extractable text (images or scans) "
                            f"and were skipped.")


def _split_blocks(raw: str) -> list:
    """Split page text into paragraph-like blocks; re-join lines broken by the PDF layout."""
    raw = raw.replace("\r", "\n")
    parts = re.split(r"\n\s*\n", raw)
    out = []
    for part in parts:
        lines = [ln.strip() for ln in part.split("\n") if ln.strip()]
        if not lines:
            continue
        buf = []
        for ln in lines:
            # a new numbered clause or bullet starts a new block
            if buf and re.match(r"^(\d+(\.\d+)*[.)]?\s|[-•*]\s|[a-z][.)]\s|\([a-z0-9]+\)\s)", ln):
                out.append(clean(" ".join(buf)))
                buf = []
            buf.append(ln)
        if buf:
            out.append(clean(" ".join(buf)))
    return [b for b in out if b]


def _read_text(data: bytes, doc: DocumentText):
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            raw = data.decode(enc)
            if enc == "utf-16" and "\x00" in raw:
                continue
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 never fails
        raise DocumentReadError("Could not decode the text file.")
    for n, block in enumerate(_split_blocks(raw), start=1):
        kind, level = "paragraph", 0
        m = re.match(r"^(#{1,6})\s+(.*)$", block)
        if m:
            kind, level, block = "heading", len(m.group(1)), m.group(2)
        doc.segments.append(Segment(f"L{n}", block, kind, level))


# --------------------------------------------------------------------------- #
# Public entry point
# --------------------------------------------------------------------------- #
def read_document(data: bytes, filename: str) -> DocumentText:
    """Read a scope document into numbered segments. Raises DocumentReadError with a user-facing message."""
    if not data:
        raise DocumentReadError("The file is empty.")
    name = filename or "document"
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    if ext == ".doc":
        raise DocumentReadError("Legacy .doc files are not supported. Open the file in Word and "
                                "Save As .docx, then upload again.")
    if ext not in SUPPORTED:
        raise DocumentReadError(f"Unsupported file type '{ext or name}'. Upload .docx, .pdf, .txt or .md.")

    doc = DocumentText(filename=name, file_type=ext.lstrip("."))
    if ext == ".docx":
        _read_docx(data, doc)
    elif ext == ".pdf":
        _read_pdf(data, doc)
    else:
        _read_text(data, doc)

    # drop exact duplicate boilerplate (repeated page headers/footers in PDFs)
    seen, kept, dups = {}, [], 0
    for s in doc.segments:
        key = s.text.lower()
        if len(key) < 120 and seen.get(key, 0) >= 2:
            dups += 1
            continue
        seen[key] = seen.get(key, 0) + 1
        kept.append(s)
    doc.segments = kept
    if dups:
        doc.warnings.append(f"{dups} repeated header/footer lines were skipped.")

    if doc.char_count < MIN_TEXT_CHARS:
        raise DocumentReadError("No usable text was found in the document.")

    if doc.char_count > MAX_CHARS:
        total, cut = 0, len(doc.segments)
        for i, s in enumerate(doc.segments):
            total += len(s.text)
            if total > MAX_CHARS:
                cut = i
                break
        last = doc.segments[cut - 1].sid if cut else "-"
        doc.segments = doc.segments[:cut]
        doc.truncated = True
        doc.warnings.append(f"The document is very long; only the text up to segment {last} "
                            f"({MAX_CHARS:,} characters) is used. Split the scope into parts for full coverage.")
    return doc
