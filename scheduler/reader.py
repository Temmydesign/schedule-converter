"""Read any schedule-type spreadsheet and turn it into a list of SourceRow.

Handles:
* .xlsx / .xlsm (openpyxl, keeps indent + bold + merged cells), .xls (xlrd), .csv
* automatic detection of the schedule sheet and of the header row
* flexible column names (Activity / Task Name / Description, Start / Start date /
  Planned Start, Dur / Duration / Dur (wks) / OD, Start Wk, WBS, Type, System /
  Phase / Area, Predecessors, ...)
* week-number schedules (Start Wk + Dur, dates taken from the timeline header)
* Gantt-only matrices (bars drawn as filled cells under a date/week header)
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from typing import Optional

from .model import SourceRow

MAX_ROWS = 3000
MAX_COLS = 250

# --------------------------------------------------------------------------- #
# Low-level sheet reading
# --------------------------------------------------------------------------- #


@dataclass
class Sheet:
    name: str
    values: list
    indent: list = field(default_factory=list)
    bold: list = field(default_factory=list)
    title_lines: list = field(default_factory=list)

    @property
    def nrows(self):
        return len(self.values)

    def cell(self, r, c):
        if r < len(self.values) and c < len(self.values[r]):
            return self.values[r][c]
        return None


def _openpyxl_sheets(data: bytes) -> list:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue
        nr = min(ws.max_row or 0, MAX_ROWS)
        nc = min(ws.max_column or 0, MAX_COLS)
        vals = [[None] * nc for _ in range(nr)]
        ind = [[0] * nc for _ in range(nr)]
        bold = [[False] * nc for _ in range(nr)]
        for r_i, row in enumerate(ws.iter_rows(min_row=1, max_row=nr, max_col=nc)):
            for c_i, cell in enumerate(row):
                vals[r_i][c_i] = cell.value
                try:
                    ind[r_i][c_i] = int(cell.alignment.indent or 0)
                    bold[r_i][c_i] = bool(cell.font and cell.font.b)
                except Exception:  # pragma: no cover - style edge cases
                    pass
        # fill vertically merged ranges (e.g. a System column merged over rows)
        for rng in ws.merged_cells.ranges:
            if rng.max_row > rng.min_row and rng.min_row - 1 < nr:
                v = vals[rng.min_row - 1][rng.min_col - 1] if rng.min_col - 1 < nc else None
                for r in range(rng.min_row - 1, min(rng.max_row, nr)):
                    for c in range(rng.min_col - 1, min(rng.max_col, nc)):
                        vals[r][c] = v
        out.append(Sheet(ws.title, vals, ind, bold))
    return out


def _xlrd_sheets(data: bytes) -> list:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    out = []
    for sh in book.sheets():
        vals = []
        for r in range(min(sh.nrows, MAX_ROWS)):
            row = []
            for c in range(min(sh.ncols, MAX_COLS)):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    except Exception:
                        pass
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                row.append(v)
            vals.append(row)
        out.append(Sheet(sh.name, vals))
    return out


def _csv_sheets(data: bytes, name: str) -> list:
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text else csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    vals = [[(c if c.strip() != "" else None) for c in r] for r in rows[:MAX_ROWS]]
    return [Sheet(name, vals)]


def read_sheets(data: bytes, filename: str) -> list:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm", "xltx", "xltm"):
        return _openpyxl_sheets(data)
    if ext == "xls":
        return _xlrd_sheets(data)
    if ext in ("csv", "txt"):
        return _csv_sheets(data, filename.rsplit(".", 1)[0])
    raise ValueError(f"Unsupported file type: .{ext}")


# --------------------------------------------------------------------------- #
# Value parsing
# --------------------------------------------------------------------------- #
_MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"


def parse_date(v, dayfirst: Optional[bool] = True) -> Optional[dt.date]:
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):
        # Excel serial date (only plausible range 1990-2100)
        if 32874 <= float(v) <= 73051:
            return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
        return None
    s = str(v).strip()
    if not s or len(s) > 40:
        return None
    s = re.sub(r"^(mon|tue|wed|thu|fri|sat|sun)[a-z]*[\s,]+", "", s, flags=re.I)
    s = re.sub(r"\s+\d{1,2}:\d{2}(:\d{2})?\s*(am|pm)?$", "", s, flags=re.I)
    if not re.search(r"\d", s):
        return None
    if not (re.search(r"\d{1,4}[/\-. ]\d{1,2}[/\-. ]\d{2,4}", s) or re.search(_MONTHS, s, re.I)):
        return None
    try:
        import pandas as pd

        if re.match(r"^\d{4}-\d{1,2}-\d{1,2}", s):
            return pd.to_datetime(s, yearfirst=True).date()
        return pd.to_datetime(s, dayfirst=bool(dayfirst)).date()
    except Exception:
        return None


def parse_number(v) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().lower()
    m = re.match(r"^([-+]?\d+(?:\.\d+)?)\s*([a-z?]*)", s)
    if not m:
        return None
    return float(m.group(1))


def duration_unit_from_text(v) -> Optional[str]:
    if isinstance(v, str):
        s = v.strip().lower()
        if re.search(r"\d\s*(w|wk|wks|week|weeks)\b", s):
            return "w"
        if re.search(r"\d\s*(d|day|days|ed|edays)\b", s):
            return "d"
        if re.search(r"\d\s*(mo|mon|month|months)\b", s):
            return "mo"
        if re.search(r"\d\s*(h|hr|hrs|hour|hours)\b", s):
            return "h"
    return None


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v).strip().lower()) if v is not None else ""


# --------------------------------------------------------------------------- #
# Header recognition
# --------------------------------------------------------------------------- #
COLUMN_PATTERNS = [
    ("start_wk", r"^(start\s*(wk|week)|(wk|week)\s*start|start\s*(wk|week)\s*(no\.?|#)?)$"),
    ("finish_wk", r"^((finish|end)\s*(wk|week)|(wk|week)\s*(finish|end))$"),
    ("preds", r"(predecessor|^preds?\.?$|depends on|dependenc|^logic$)"),
    ("succs", r"(successor)"),
    ("id", r"^(id|task id|activity id|act\.? id|activity code|unique id|uid|code|activity no\.?|task no\.?)$"),
    ("type", r"^(type|activity type|task type|act\.? type|category|item type)$"),
    ("wbs", r"^(wbs|wbs (code|id|no\.?|number|ref)|ref\.?|ref\.? no\.?|item( no\.?)?|s/?n|s\.n\.?|no\.?|#|outline( number| no\.?)?|level)$"),
    ("group", r"^(system|sub-?system|phase|area|discipline|group|stage|work ?stream|package|work package|unit|section|department|location|facility|asset|module|zone|contract|lot)$"),
    ("start", r"^((planned|plan|early|baseline|bl|scheduled|target|forecast|actual)\s+)?(start|begin)(\s*date)?$|^start date\b|^from$"),
    ("finish", r"^((planned|plan|early|baseline|bl|scheduled|target|forecast|actual)\s+)?(finish|end|completion|complete|due)(\s*date)?$|^finish date\b|^to$"),
    ("duration", r"^(dur\.?|duration|original duration|orig\.? dur\.?|od|planned duration|days|weeks|wks)(\s*\((wks?|weeks?|days?|d|w)\))?$|^dur(ation)?\s*[\(\[]"),
    ("name", r"^(activity|activities|task|tasks|activity name|task name|name|description|activity description|task description|deliverable|deliverables|scope|work item|item description|title|activity title|milestone|milestones|work description)$"),
]


def classify_header(text) -> Optional[str]:
    s = _norm(text)
    if not s or len(s) > 60:
        return None
    s = s.rstrip(":")
    for key, pat in COLUMN_PATTERNS:
        if re.search(pat, s):
            return key
    return None


@dataclass
class HeaderInfo:
    row: int
    cols: dict
    score: float
    timeline: dict = field(default_factory=dict)   # col -> period start date
    timeline_step: int = 7
    duration_header: str = ""


def _find_timeline(sheet: Sheet, header_row: int, used_cols: set):
    """Detect a Gantt timeline (dates or week numbers) in or above the header row."""
    best = {}
    step = 7
    for r in range(max(0, header_row - 3), min(sheet.nrows, header_row + 1)):
        dates = {}
        for c, v in enumerate(sheet.values[r]):
            if c in used_cols:
                continue
            if isinstance(v, dt.datetime):
                d = v.date()
            elif isinstance(v, dt.date):
                d = v
            elif isinstance(v, str):
                d = parse_date(v)
            else:
                d = None
            if d:
                dates[c] = d
        if len(dates) >= 4 and len(dates) > len(best):
            cols = sorted(dates)
            diffs = [(dates[b] - dates[a]).days for a, b in zip(cols, cols[1:])]
            if diffs and all(x > 0 for x in diffs):
                best = dates
                step = max(set(diffs), key=diffs.count)
    return best, step


def _week_number_map(sheet: Sheet, header_row: int, timeline: dict):
    """Map week numbers (1..N in the header row under the timeline) to dates."""
    wk = {}
    for c, d in timeline.items():
        v = sheet.cell(header_row, c)
        n = parse_number(v)
        if n is not None and float(n).is_integer():
            wk[int(n)] = d
    return wk


def detect_header(sheet: Sheet) -> Optional[HeaderInfo]:
    best = None
    for r in range(min(sheet.nrows, 40)):
        cols = {}
        for c, v in enumerate(sheet.values[r]):
            if not isinstance(v, str):
                continue
            key = classify_header(v)
            if key and key not in cols:
                cols[key] = c
        if "name" not in cols:
            continue
        has_dates = ("start" in cols and ("finish" in cols or "duration" in cols)) or \
                    ("finish" in cols and "duration" in cols)
        has_weeks = "start_wk" in cols and ("duration" in cols or "finish_wk" in cols)
        score = len(cols) + (5 if has_dates else 0) + (3 if has_weeks else 0)
        # count data rows beneath
        n_data = 0
        for rr in range(r + 1, min(sheet.nrows, r + 400)):
            v = sheet.cell(rr, cols["name"])
            if v not in (None, "") and str(v).strip():
                n_data += 1
        if n_data < 1:
            continue
        info = HeaderInfo(r, cols, score + min(n_data, 200) / 50)
        info.timeline, info.timeline_step = _find_timeline(sheet, r, set(cols.values()))
        if not (has_dates or has_weeks) and len(info.timeline) < 4:
            continue  # nothing to schedule from
        if info.timeline:
            info.score += 1
        if "duration" in cols:
            info.duration_header = str(sheet.cell(r, cols["duration"]))
        if best is None or info.score > best.score:
            best = info
    return best


_SHEET_NAME_HINTS = re.compile(r"(schedul|programme|program|gantt|l[1-5]\b|level\s*[1-5]|plan|timeline|baseline|activities)", re.I)


def rank_sheets(sheets: list) -> list:
    """Return [(sheet, header_info, score)] sorted best first (only plausible sheets)."""
    ranked = []
    for sh in sheets:
        h = detect_header(sh)
        if not h:
            continue
        score = h.score + (4 if _SHEET_NAME_HINTS.search(sh.name) else 0)
        ranked.append((sh, h, score))
    ranked.sort(key=lambda x: -x[2])
    return ranked


# --------------------------------------------------------------------------- #
# Row extraction
# --------------------------------------------------------------------------- #
GANTT_MILESTONE_MARKS = {"◆", "♦", "◇", "★", "m", "ms", "x◆", "▲", "▼", "⬥", "●"}


def _is_milestone_type(t) -> bool:
    if not t:
        return False
    s = _norm(t)
    return bool(re.search(r"(mile|mstone|^ms$|^mlst|gate|key ?date|^m$|decision)", s))


def _is_summary_type(t) -> bool:
    if not t:
        return False
    s = _norm(t)
    return bool(re.search(r"^(sum|summary|header|heading|hammock|wbs|phase|section|group)$", s))


def extract_rows(sheet: Sheet, h: HeaderInfo, project_start: Optional[dt.date] = None,
                 dayfirst: Optional[bool] = True, days_per_week: int = 5):
    """Return (rows, ignored, notes)."""
    cols = h.cols
    notes = []
    ignored = []
    rows: list = []

    week_map = _week_number_map(sheet, h.row, h.timeline) if h.timeline else {}
    week1 = week_map.get(1)
    if not week1 and h.timeline:
        week1 = min(h.timeline.values())
    base_for_weeks = week1 or project_start

    # duration unit from header
    dh = _norm(h.duration_header)
    header_unit = "w" if re.search(r"(wk|week)", dh) else ("d" if re.search(r"(day|\(d\))", dh) else None)

    timeline_cols = sorted(h.timeline)

    def gantt_span(r):
        marked = [c for c in timeline_cols if sheet.cell(r, c) not in (None, "", 0)]
        if not marked:
            return None, None, False
        first, last = marked[0], marked[-1]
        s = h.timeline[first]
        f = h.timeline[last] + dt.timedelta(days=h.timeline_step - 1)
        if h.timeline_step == 7:
            f = h.timeline[last] + dt.timedelta(days=days_per_week - 1)
        mark = str(sheet.cell(r, first)).strip().lower()
        is_ms = len(marked) == 1 and mark in GANTT_MILESTONE_MARKS
        return s, f, is_ms

    for r in range(h.row + 1, sheet.nrows):
        raw_name = sheet.cell(r, cols["name"])
        name = re.sub(r"\s+", " ", str(raw_name)).strip() if raw_name not in (None, "") else ""
        if not name:
            continue
        get = lambda k: sheet.cell(r, cols[k]) if k in cols else None  # noqa: E731
        start = parse_date(get("start"), dayfirst)
        finish = parse_date(get("finish"), dayfirst)
        dur_raw = get("duration")
        dur = parse_number(dur_raw)
        unit = duration_unit_from_text(dur_raw) or header_unit
        swk = parse_number(get("start_wk"))
        fwk = parse_number(get("finish_wk"))
        typ = get("type")
        typ = str(typ).strip() if typ not in (None, "") else None

        # week-number schedules -------------------------------------------- #
        if start is None and swk is not None and base_for_weeks:
            start = base_for_weeks + dt.timedelta(weeks=int(swk) - 1)
            if unit is None:
                unit = "w"
            if fwk is not None:
                finish = base_for_weeks + dt.timedelta(weeks=int(fwk) - 1, days=days_per_week - 1)
            elif dur is not None:
                finish = start + dt.timedelta(weeks=max(int(dur), 1) - 1, days=days_per_week - 1)

        gantt_ms = False
        if start is None and finish is None and timeline_cols:
            start, finish, gantt_ms = gantt_span(r)

        idx_cell = lambda k: sheet.cell(r, cols[k]) if k in cols else None  # noqa: E731
        indent = sheet.indent[r][cols["name"]] if sheet.indent and cols["name"] < len(sheet.indent[r]) else 0
        bold = sheet.bold[r][cols["name"]] if sheet.bold and cols["name"] < len(sheet.bold[r]) else False

        wbs = idx_cell("wbs")
        wbs = _clean_code(wbs)
        sid = _clean_code(idx_cell("id"))
        grp = idx_cell("group")
        grp = str(grp).strip() if grp not in (None, "") else None
        preds = idx_cell("preds")
        preds = str(preds).strip() if preds not in (None, "") else None
        if isinstance(idx_cell("preds"), float) and float(idx_cell("preds")).is_integer():
            preds = str(int(idx_cell("preds")))

        row = SourceRow(
            excel_row=r + 1, name=name, start=start, finish=finish, duration=dur,
            dur_unit=unit or "", wbs=wbs, type=typ, group=grp, src_id=sid, preds=preds,
            indent=indent, bold=bold,
        )
        row.is_milestone = _is_milestone_type(typ) or (dur is not None and dur == 0) or gantt_ms
        row.is_summary = _is_summary_type(typ)
        rows.append(row)

    # ---------------------------------------------------------------- #
    # Classify dateless rows: section headers vs trailing notes/legend
    # ---------------------------------------------------------------- #
    last_dated = max((i for i, x in enumerate(rows) if x.start or x.finish), default=-1)
    kept = []
    for i, x in enumerate(rows):
        dated = bool(x.start or x.finish)
        if dated:
            kept.append(x)
            continue
        if i < last_dated and len(x.name) <= 150 and not x.name.lower().startswith(("note", "legend")):
            x.is_summary = True          # section heading
            kept.append(x)
        else:
            ignored.append((x.excel_row, x.name, "no dates / notes or legend row"))

    # forward-fill group column when it is only filled at block starts
    if "group" in cols:
        filled = sum(1 for x in kept if x.group)
        if 0 < filled < len(kept):
            cur = None
            for x in kept:
                if x.group:
                    cur = x.group
                elif not x.is_summary:
                    x.group = cur

    # fill missing start/finish from duration
    for x in kept:
        if x.is_summary:
            continue
        if x.start and not x.finish and x.duration is not None:
            x.finish = _add_working(x.start, x.duration, x.dur_unit or "d", days_per_week)
        if x.finish and not x.start and x.duration is not None:
            x.start = _add_working(x.finish, -x.duration, x.dur_unit or "d", days_per_week)
        if x.start and x.finish and x.finish < x.start:
            notes.append(f"Row {x.excel_row}: finish before start - swapped.")
            x.start, x.finish = x.finish, x.start
    return kept, ignored, notes


def _add_working(d: dt.date, dur: float, unit: str, days_per_week: int = 5) -> dt.date:
    import numpy as np

    mask = [1] * days_per_week + [0] * (7 - days_per_week)
    days = int(round(dur * (days_per_week if unit == "w" else 1)))
    if days == 0:
        return d
    step = days - 1 if days > 0 else days + 1
    return np.busday_offset(np.datetime64(d), step, roll="forward", weekmask=mask).astype(dt.date)


def _clean_code(v) -> Optional[str]:
    if v in (None, ""):
        return None
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(v)
    s = str(v).strip()
    return s or None


# --------------------------------------------------------------------------- #
# Project title suggestion
# --------------------------------------------------------------------------- #
def suggest_project_name(sheet: Sheet, header_row: int, filename: str) -> str:
    lines = []
    for r in range(0, header_row):
        for v in sheet.values[r]:
            if isinstance(v, str) and len(v.strip()) > 8:
                lines.append(v.strip())
                break
    base, sub = None, None
    for line in lines:
        if "|" in line:
            for seg in [s.strip() for s in line.split("|")]:
                low = seg.lower()
                if not seg or low.startswith("prepared") or re.match(r"^[\w./-]+$", seg):
                    continue
                m = re.match(r"^[A-Z0-9]{2,6}\s+-\s+(.+?)(\s*\(.*\))?$", seg)
                if m and base:
                    sub = m.group(1).strip()
                    continue
                if base is None:
                    base = re.sub(r"\s*\(.*?\)", "", seg)
                    base = re.split(r"\s+-\s+", base)[0].strip()
    sheet_title = re.sub(r"\bL([1-5])\b", r"Level \1", sheet.name).strip()
    if base:
        parts = [base] + ([sub] if sub else []) + [sheet_title]
        return " - ".join(parts)
    if lines:
        return re.sub(r"\s+", " ", lines[0])[:120]
    return re.sub(r"[_]+", " ", filename.rsplit(".", 1)[0]).strip()
