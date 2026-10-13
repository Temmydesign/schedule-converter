"""Approved scope register -> verified CPM schedule (deterministic; no AI in this module).

Pipeline
--------
1. ``validate_register``   Review-gate checks on the planner-edited register (blocking errors + warnings).
2. ``compute_dates``       Forward pass on working days from durations and FS/SS/FF logic.
3. ``standard_workbook``   Writes the firm's *standard layout* workbook (the same one the Excel converter
                           already reads: ID, WBS, Activity, Type, Start, Finish, Duration, Predecessors).
4. ``build_schedule``      Runs the **unchanged** proven engine (``engine.convert``) on that workbook, so the
                           outline, firm policy, logic handling, CPM verification and every exporter are the
                           same code paths as the Excel converter. The engine's independent forward pass must
                           reproduce the register dates exactly (zero-day variance) or the build fails.
5. ``tie_completion_chain`` Firm-policy close-out for document schedules: every summary's
                           "<summary> completed" milestone drives its parent's, so the network has one start
                           and one finish (no dangling ends).

The LLM never sets a date, a float value or a relationship that the planner has not approved.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import re
from collections import defaultdict, deque
from dataclasses import dataclass, field
from typing import Optional

import pandas as pd

from . import calendar_detect, cpm, engine, exporters
from .model import Link, WorkCalendar

# --------------------------------------------------------------------------- #
# Register layout (shared with extractor.py and app.py)
# --------------------------------------------------------------------------- #
REGISTER_COLUMNS = ["Include", "Ref", "WBS", "Activity", "Type", "Duration (days)", "Predecessors",
                    "Basis", "Confidence", "Citation", "Source", "Evidence", "Assumption", "Planner note"]
TYPES = ["Summary", "Task", "Milestone"]
CITATION_OK = {"Verified", "Relocated"}

LOW_CONFIDENCE = 0.60
DCMA_HIGH_DURATION = 44          # working days (DCMA check 8)
DCMA_PCT = 5.0                   # % threshold for lags / logic / high duration
DCMA_FS_MIN_PCT = 90.0           # % of FS relationships (DCMA check 4)

_WBS_RE = re.compile(r"^\d+(\.\d+)*$")
_PRED_RE = re.compile(
    r"^(?P<ref>[A-Za-z0-9][A-Za-z0-9._\-/]*?)\s*(?P<type>FS|SS|FF|SF)?\s*"
    r"(?:(?P<lag>[+-]\s*\d+(?:\.\d+)?)\s*(?:d|day|days|wd)?)?$", re.I)


class BuildError(ValueError):
    """Raised when an approved register cannot produce a verified schedule; carries the reasons."""

    def __init__(self, message: str, details: Optional[list] = None):
        super().__init__(message)
        self.details = details or []


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #
def wbs_key(code: str) -> tuple:
    try:
        return tuple(int(p) for p in str(code).split("."))
    except ValueError:
        return (10 ** 9, str(code))


def wbs_parent(code: str) -> Optional[str]:
    return code.rsplit(".", 1)[0] if "." in code else None


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and pd.isna(v):
        return ""
    return str(v).strip()


def _num(v) -> Optional[float]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if pd.isna(f) else f


def _bool(v) -> bool:
    if isinstance(v, str):
        return v.strip().lower() in ("true", "yes", "1", "y", "x")
    try:
        return bool(v) and not pd.isna(v)
    except TypeError:
        return bool(v)


def parse_predecessors(text: str):
    """'A010, A020 SS+5d, A030FF' -> ([(ref, type, lag_days)], [bad tokens])."""
    out, bad = [], []
    for tok in re.split(r"[,;\n]+", _txt(text)):
        tok = tok.strip()
        if not tok or tok.lower() in ("-", "none", "nil", "n/a"):
            continue
        m = _PRED_RE.match(tok)
        if not m:
            bad.append(tok)
            continue
        lag = m.group("lag")
        lag_v = float(re.sub(r"\s+", "", lag)) if lag else 0.0
        out.append((m.group("ref").upper(), (m.group("type") or "FS").upper(), lag_v))
    return out, bad


def format_predecessors(deps) -> str:
    """[(ref, type, lag)] -> 'A010, A020 SS+5d'."""
    parts = []
    for ref, typ, lag in deps:
        lag = int(round(lag or 0))
        if typ == "FS" and lag == 0:
            parts.append(ref)
        else:
            parts.append(f"{ref} {typ}{'' if lag == 0 else f'{lag:+d}d'}")
    return ", ".join(parts)


# --------------------------------------------------------------------------- #
# 1. Review-gate validation
# --------------------------------------------------------------------------- #
@dataclass
class Validation:
    errors: list = field(default_factory=list)        # (ref, message)  -> block approval
    warnings: list = field(default_factory=list)      # (ref, message)  -> planner judgement
    flags: dict = field(default_factory=dict)         # ref -> [short flag]
    stats: dict = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors

    def flag(self, ref, text):
        self.flags.setdefault(ref, []).append(text)


@dataclass(eq=False)
class Item:
    ref: str
    wbs: str
    name: str
    type: str
    dur: int
    preds: list                     # [(ref, type, lag_days:int)]
    row: dict
    es: Optional[int] = None
    ef: Optional[int] = None


def register_items(df: pd.DataFrame) -> list:
    """Included rows as Items, in WBS order (no validation)."""
    items = []
    for _, r in df.iterrows():
        if not _bool(r.get("Include", True)):
            continue
        typ = _txt(r.get("Type")).title() or "Task"
        dur = _num(r.get("Duration (days)"))
        deps, _ = parse_predecessors(r.get("Predecessors"))
        items.append(Item(ref=_txt(r.get("Ref")).upper(), wbs=_txt(r.get("WBS")).rstrip("."),
                          name=_txt(r.get("Activity")), type=typ,
                          dur=0 if typ != "Task" else int(round(dur or 0)),
                          preds=[(a, b, int(round(c))) for a, b, c in deps], row=dict(r)))
    items.sort(key=lambda it: wbs_key(it.wbs))
    return items


def validate_register(df: pd.DataFrame, max_lead_days: int = 0) -> Validation:
    v = Validation()
    if df is None or df.empty:
        v.errors.append(("", "The register is empty."))
        return v

    all_rows = {}
    for _, r in df.iterrows():
        ref = _txt(r.get("Ref")).upper()
        if ref:
            all_rows.setdefault(ref, r)
    included = [r for _, r in df.iterrows() if _bool(r.get("Include", True))]
    excluded_refs = {_txt(r.get("Ref")).upper() for _, r in df.iterrows() if not _bool(r.get("Include", True))}
    if not included:
        v.errors.append(("", "No rows are included. Tick 'Include' for the scope to schedule."))
        return v

    refs, wbs_seen = {}, {}
    for r in included:
        ref = _txt(r.get("Ref")).upper()
        name = _txt(r.get("Activity"))
        wbs = _txt(r.get("WBS")).rstrip(".")
        typ = _txt(r.get("Type")).title()
        key = ref or f"(row '{name[:30]}')"
        if not ref:
            v.errors.append((key, "Ref is blank. Give every row a unique reference (e.g. A150)."))
        elif ref in refs:
            v.errors.append((ref, f"Ref {ref} is used more than once."))
        else:
            refs[ref] = r
        if not name:
            v.errors.append((key, "Activity name is blank."))
        if typ not in TYPES:
            v.errors.append((key, f"Type '{_txt(r.get('Type'))}' must be Summary, Task or Milestone."))
        if not _WBS_RE.match(wbs):
            v.errors.append((key, f"WBS '{wbs}' is not a numeric outline code (e.g. 2.3.1)."))
        elif wbs in wbs_seen:
            v.errors.append((key, f"WBS {wbs} is also used by {wbs_seen[wbs]}."))
        else:
            wbs_seen[wbs] = ref or name[:30]

        dur = _num(r.get("Duration (days)"))
        if typ == "Task":
            if dur is None or dur <= 0:
                v.errors.append((key, "Task has no duration. Enter working days (or change Type to Milestone)."))
                v.flag(ref, "No duration")
            elif abs(dur - round(dur)) > 1e-9:
                v.warnings.append((key, f"Duration {dur:g} rounded to {round(dur)} working days."))
            elif dur > DCMA_HIGH_DURATION:
                v.flag(ref, "High duration")
        elif typ == "Milestone" and dur not in (None, 0):
            v.warnings.append((key, f"Milestone duration {dur:g} ignored (milestones are 0 days)."))

        conf = _num(r.get("Confidence"))
        if conf is not None and conf < LOW_CONFIDENCE:
            v.warnings.append((key, f"Low extraction confidence ({conf:.0%}). Check against the source."))
            v.flag(ref, "Low confidence")
        cit = _txt(r.get("Citation"))
        if cit not in CITATION_OK and not _txt(r.get("Planner note")):
            why = ("Row added by the planner" if cit in ("", "Planner")
                   else "Evidence quote not found in the document")
            v.errors.append((key, f"{why}: add a Planner note to justify it (traceability)."))
            v.flag(ref, "Note required")
        if _txt(r.get("Basis")).lower() == "estimated":
            v.flag(ref, "Estimated duration")

    # hierarchy -------------------------------------------------------------- #
    by_wbs = {_txt(r.get("WBS")).rstrip("."): r for r in included}
    has_children = defaultdict(int)
    for w in by_wbs:
        p = wbs_parent(w)
        if p is not None:
            if p not in by_wbs:
                v.errors.append((_txt(by_wbs[w].get("Ref")).upper(),
                                 f"WBS {w} has no parent row {p}. Add the summary or renumber."))
            else:
                has_children[p] += 1
    for w, r in by_wbs.items():
        ref = _txt(r.get("Ref")).upper()
        typ = _txt(r.get("Type")).title()
        if has_children.get(w) and typ != "Summary":
            v.errors.append((ref, f"WBS {w} has child rows, so its Type must be Summary."))
        if typ == "Summary" and not has_children.get(w):
            v.errors.append((ref, "Summary with no activities under it. Add activities or change the Type."))

    # logic ------------------------------------------------------------------ #
    links, fs = 0, 0
    lags = 0
    succ_count = defaultdict(int)
    graph = defaultdict(list)
    types = {ref: _txt(r.get("Type")).title() for ref, r in refs.items()}
    for ref, r in refs.items():
        deps, bad = parse_predecessors(r.get("Predecessors"))
        for b in bad:
            v.errors.append((ref, f"Predecessor '{b}' is not understood. Use e.g. A010, A020 SS+5d, A030 FF."))
        if deps and types.get(ref) == "Summary":
            v.errors.append((ref, "Logic on a summary row is not allowed (DCMA). Put it on the first activity."))
        for p, typ, lag in deps:
            if p == ref:
                v.errors.append((ref, f"Lists itself ({p}) as a predecessor."))
                continue
            if p not in refs:
                why = "is excluded" if p in excluded_refs else ("is not in the register" if p not in all_rows
                                                                 else "is not included")
                v.errors.append((ref, f"Predecessor {p} {why}."))
                continue
            if typ == "SF":
                v.errors.append((ref, f"Start-to-Finish link from {p}: not allowed (DCMA check 4)."))
            if lag < 0 and -lag > max_lead_days:
                v.errors.append((ref, f"Lead (negative lag) of {lag:g}d from {p}: leads are not allowed "
                                      f"(DCMA check 2). Use SS with a positive lag instead."))
            if types.get(p) == "Summary" and typ != "FS":
                v.errors.append((ref, f"Link from summary {p} must be FS (it follows the last activities of {p})."))
            links += 1
            fs += typ == "FS"
            lags += lag > 0
            succ_count[p] += 1
            graph[p].append(ref)

    # loops ------------------------------------------------------------------ #
    cyc = find_cycle(list(refs), graph)
    if cyc:
        v.errors.append((cyc[0], "Circular logic: " + " -> ".join(cyc + [cyc[0]])))

    # planning-quality warnings (DCMA style) --------------------------------- #
    tasks = [ref for ref, t in types.items() if t == "Task"]
    if links:
        fs_pct = 100.0 * fs / links
        lag_pct = 100.0 * lags / links
        if fs_pct < DCMA_FS_MIN_PCT:
            v.warnings.append(("", f"Only {fs_pct:.0f}% of links are FS (DCMA target >= {DCMA_FS_MIN_PCT:.0f}%)."))
        if lag_pct > DCMA_PCT:
            v.warnings.append(("", f"{lag_pct:.0f}% of links carry a lag (DCMA target <= {DCMA_PCT:.0f}%)."))
    starts = [ref for ref in refs if types.get(ref) != "Summary"
              and not parse_predecessors(refs[ref].get("Predecessors"))[0]]
    if len(starts) > 1:
        v.warnings.append(("", f"{len(starts)} activities have no predecessor; firm policy will link them SS to "
                               f"the project's first activity. Check they really start on day one: "
                               + ", ".join(starts[:10]) + ("…" if len(starts) > 10 else "")))
    first = next((r for r in sorted(refs.values(), key=lambda x: wbs_key(_txt(x.get("WBS"))))
                  if _txt(r.get("Type")).title() != "Summary"), None)
    if first is not None and _txt(first.get("Type")).title() != "Milestone":
        v.warnings.append(("", "The schedule does not open with a start milestone (e.g. Contract Award / NTP)."))
    high = [ref for ref in tasks if (_num(refs[ref].get("Duration (days)")) or 0) > DCMA_HIGH_DURATION]
    if tasks and 100.0 * len(high) / len(tasks) > DCMA_PCT:
        v.warnings.append(("", f"{len(high)} tasks exceed {DCMA_HIGH_DURATION} working days "
                               f"(DCMA check 8 target <= {DCMA_PCT:.0f}% of tasks). Consider splitting them."))

    v.stats = {"included": len(included), "excluded": len(df) - len(included),
               "summaries": sum(1 for t in types.values() if t == "Summary"),
               "tasks": len(tasks), "milestones": sum(1 for t in types.values() if t == "Milestone"),
               "links": links, "fs_pct": (100.0 * fs / links) if links else 0.0,
               "estimated": sum(1 for r in included if _txt(r.get("Basis")).lower() == "estimated"),
               "unverified": sum(1 for r in included if _txt(r.get("Citation")) not in CITATION_OK)}
    return v


def find_cycle(nodes: list, graph: dict) -> Optional[list]:
    """Return one cycle as a list of node ids, or None (iterative DFS, colour marking)."""
    color = {n: 0 for n in nodes}
    parent = {}
    for root in nodes:
        if color[root]:
            continue
        stack = [(root, iter(graph.get(root, ())))]
        color[root] = 1
        while stack:
            node, it = stack[-1]
            nxt = next(it, None)
            if nxt is None:
                color[node] = 2
                stack.pop()
                continue
            if nxt not in color:
                continue
            if color[nxt] == 0:
                color[nxt] = 1
                parent[nxt] = node
                stack.append((nxt, iter(graph.get(nxt, ()))))
            elif color[nxt] == 1:
                cyc, cur = [nxt], node
                while cur != nxt:
                    cyc.append(cur)
                    cur = parent[cur]
                cyc.reverse()
                return [cyc[-1]] + cyc[:-1] if cyc else cyc
    return None


# --------------------------------------------------------------------------- #
# 2. Deterministic dates
# --------------------------------------------------------------------------- #
def compute_dates(items: list) -> None:
    """Forward pass on working-day boundaries (same time model as model.py / cpm.py).

    A predecessor that is a summary is resolved to all activities under it (FS = after the
    last of them finishes), which is what the engine's "<summary> completed" milestone does.
    """
    by_ref = {it.ref: it for it in items}
    leaves_of = defaultdict(list)
    for it in items:
        if it.type == "Summary":
            continue
        p = wbs_parent(it.wbs)
        while p is not None:
            leaves_of[p].append(it)
            p = wbs_parent(p)
    def expand(ref):
        it = by_ref[ref]
        return leaves_of.get(it.wbs, []) if it.type == "Summary" else [it]

    sched = [it for it in items if it.type != "Summary"]
    deps = {it: [] for it in sched}
    for it in sched:
        for p, typ, lag in it.preds:
            for leaf in expand(p):
                deps[it].append((leaf, typ, lag))
    indeg = {it: len(deps[it]) for it in sched}
    succ = defaultdict(list)
    for it, ds in deps.items():
        for leaf, _, _ in ds:
            succ[leaf].append(it)
    q = deque(it for it in sched if indeg[it] == 0)
    done = 0
    while q:
        it = q.popleft()
        done += 1
        es = 0
        for leaf, typ, lag in deps[it]:
            if typ == "FS":
                v = leaf.ef + lag
            elif typ == "SS":
                v = leaf.es + lag
            elif typ == "FF":
                v = leaf.ef + lag - it.dur
            else:
                v = leaf.es + lag - it.dur
            es = max(es, v)
        it.es, it.ef = es, es + it.dur
        for s in succ[it]:
            indeg[s] -= 1
            if indeg[s] == 0:
                q.append(s)
    if done != len(sched):
        raise BuildError("Circular logic in the approved register.")
    for it in sorted((x for x in items if x.type == "Summary"), key=lambda x: -len(x.wbs.split("."))):
        kids = leaves_of.get(it.wbs, [])
        it.es = min(k.es for k in kids)
        it.ef = max(k.ef for k in kids)


def _display(it: Item, cal: WorkCalendar):
    """Start/finish dates as MS Project will show them (cf. cpm.display_dates)."""
    if it.type == "Milestone":
        d = cal.date(it.es - 1) if it.es > 0 else cal.date(0)
        return d, d
    if it.type == "Summary":
        return None, None
    return cal.date(it.es), cal.date(it.ef - 1)


# --------------------------------------------------------------------------- #
# 3 + 4. Standard workbook -> proven engine
# --------------------------------------------------------------------------- #
@dataclass
class BuildSettings:
    project_name: str
    start: dt.date
    work_week: int = 5
    hours_per_day: float = 8.0
    holidays: list = field(default_factory=list)
    calendar_name: Optional[str] = None
    policy: dict = field(default_factory=dict)       # overrides on firm_policy.json (from the sidebar)
    close_out: bool = True                           # tie summary completions into one finish (no open ends)


@dataclass
class BuildOutput:
    res: object                                      # engine.Result
    workbook: bytes                                  # standard-layout workbook fed to the engine
    log: list                                        # human-readable build steps


def standard_rows(items: list, cal: WorkCalendar) -> list:
    ids = {it.ref: i for i, it in enumerate(items, start=1)}
    leaves_of = defaultdict(list)
    for it in items:
        if it.type != "Summary":
            p = wbs_parent(it.wbs)
            while p is not None:
                leaves_of[p].append(it)
                p = wbs_parent(p)
    by_ref = {it.ref: it for it in items}
    internal_succ = defaultdict(set)          # leaf -> successors inside each summary
    for it in items:
        for p, _, _ in it.preds:
            internal_succ[p].add(it.ref)

    def _resolve(ref):
        """A summary predecessor = its terminal activities (+ any leaf finishing last)."""
        P = by_ref[ref]
        if P.type != "Summary":
            return [ref]
        leaves = leaves_of[P.wbs]
        inside = {x.ref for x in leaves}
        last = max(x.ef for x in leaves)
        keep = [x.ref for x in leaves if not (internal_succ[x.ref] & inside) or x.ef == last]
        return keep

    rows = []
    for it in items:
        if it.type == "Summary":
            spans = [_display(k, cal) for k in leaves_of[it.wbs]]
            s, f = min(x[0] for x in spans), max(x[1] for x in spans)
        else:
            s, f = _display(it, cal)
        preds = []
        for p, typ, lag in it.preds:
            lag_t = "" if lag == 0 else f"{lag:+d}d"
            for q in _resolve(p):
                tok = f"{ids[q]}{'' if (typ == 'FS' and not lag_t) else typ}{lag_t}"
                if tok not in preds:
                    preds.append(tok)
        if it.type == "Milestone" and it.es > 0:
            s = None        # finish-only milestone = end of that day (engine: set_targets), never the day start
        rows.append({"ID": ids[it.ref], "WBS": it.wbs, "Activity": it.name, "Type": it.type,
                     "Start": s, "Finish": f,
                     "Duration (days)": None if (it.type == "Summary" or s is None) else it.dur,
                     "Predecessors": ", ".join(preds)})
    return rows


def _check_precedence(items: list) -> list:
    """The engine drops any link whose predecessor starts after its successor; catch that up front."""
    by_ref = {it.ref: it for it in items}
    problems = []
    for it in items:
        if it.type == "Summary":
            continue
        for p, typ, lag in it.preds:
            P = by_ref[p]
            if P.type != "Summary" and P.es > it.es:
                problems.append((it.ref, f"{p} {typ}{'' if not lag else f'{lag:+d}d'} lets {it.ref} start "
                                         f"{P.es - it.es} working day(s) before {p} starts. Use FS or SS, or "
                                         f"shorten {it.ref}."))
    return problems


def build_schedule(df: pd.DataFrame, settings: BuildSettings) -> BuildOutput:
    """Approved register -> verified engine Result. Raises BuildError if anything does not verify."""
    log = []
    v = validate_register(df, max_lead_days=0)
    if not v.ok:
        raise BuildError("The register has blocking errors.", [f"{r}: {m}" if r else m for r, m in v.errors])
    items = register_items(df)
    working_days = calendar_detect.week_days(settings.work_week)
    cal = WorkCalendar(settings.start, working_days, settings.holidays, settings.hours_per_day)
    compute_dates(items)
    problems = _check_precedence(items)
    if problems:
        raise BuildError("Some relationships cannot be scheduled as written.", [f"{r}: {m}" for r, m in problems])
    log.append(f"Forward pass: {len(items)} rows, finish on {cal.date(max(i.ef for i in items) - 1):%d %b %Y} "
               f"({max(i.ef for i in items)} working days).")

    rows = standard_rows(items, cal)
    title = settings.project_name.strip() or "Project"
    wb = exporters._write_standard(rows, [("Project", title), ("Source", "Approved scope register"),
                                          ("Start date", f"{cal.date(0):%d/%m/%Y}")], title)
    log.append("Standard-layout workbook written (same format the Excel converter reads).")

    policy = dict(settings.policy or {})
    if settings.close_out:
        policy["completion_milestone_for_project_title"] = True
        policy["close_open_ends"] = True
    multilevel = any("." in it.wbs for it in items)
    opts = engine.Options(project_name=title, calendar_name=settings.calendar_name, dayfirst=True,
                          holidays=list(settings.holidays), use_source_logic=True, fit_source_logic=True,
                          hierarchy_mode="wbs" if multilevel else None, policy=policy,
                          work_week=settings.work_week, hours_per_day=settings.hours_per_day)
    res = engine.convert(wb, f"{_safe(title)}.xlsx", opts)
    log.append("Proven engine run: outline, firm policy, logic import, CPM forward/backward pass.")

    fitted = [i for i in res.issues if i[0] == "info" and "Link from" in str(i[3])]
    if settings.close_out:
        added = tie_completion_chain(res)
        if added:
            log.append(f"Close-out: {added} summary completion milestones chained to their parent's "
                       f"(single project finish).")

    st = res.stats
    removed = [i for i in res.issues if i[0] == "warning" and ("removed" in str(i[3]) or "not found" in str(i[3]))]
    if st["date_mismatches"] or not st["finish_matches"] or removed:
        detail = [f"{i[1]} {i[2]}: {i[3]}" for i in res.issues if i[0] in ("error",) or i in removed][:30]
        raise BuildError("Zero-day variance check failed: the engine did not reproduce the register dates.",
                         detail)
    log.append(f"Verified: {st['activities'] + st['milestones']} activities and milestones match the register "
               f"dates exactly (0-day variance); {len(fitted)} link(s) adjusted by the engine.")
    return BuildOutput(res=res, workbook=wb, log=log)


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")[:60] or "Schedule"


# --------------------------------------------------------------------------- #
# 5. Close-out chain (applied after the engine; mirrors engine.apply_edited_predecessors)
# --------------------------------------------------------------------------- #
def tie_completion_chain(res) -> int:
    """Link every summary's completion milestone FS to its parent's completion milestone.

    Dates cannot move: a parent completion already sits at or after every child's finish.
    Result: one project finish milestone and no completion milestones left without a successor.
    """
    added = 0
    for n in res.nodes:
        if n.kind != "completion" or n.parent is None or n.parent.parent is None:
            continue
        grand = n.parent.parent
        target = next((c for c in grand.children if c.kind == "completion"), None)
        if target is None or any(l.pred is n for l in target.preds):
            continue
        target.preds.append(Link(n, "FS", 0, "policy: close-out chain"))
        added += 1
    if added:
        for n in res.nodes:
            n.es = n.ef = None
        cpm.forward_pass(res.nodes, 0)
        res.floats = cpm.backward_pass(res.nodes)
        res.issues, res.stats = engine.quality_checks(res.nodes, res.cal, res.out_cal, res.rows, None,
                                                      res.anchor, res.floats, [], res.policy)
    return added


# --------------------------------------------------------------------------- #
# Approval record (audit trail)
# --------------------------------------------------------------------------- #
AUDIT_COLUMNS = ["Include", "Ref", "WBS", "Activity", "Type", "Duration (days)", "Predecessors", "Basis",
                 "Confidence", "Citation", "Source", "Evidence", "Assumption", "Planner note"]


def canonical_register(df: pd.DataFrame) -> list:
    out = []
    for _, r in df.iterrows():
        rec = {}
        for c in AUDIT_COLUMNS:
            v = r.get(c)
            if c == "Include":
                rec[c] = _bool(v)
            elif c in ("Duration (days)", "Confidence"):
                n = _num(v)
                rec[c] = None if n is None else round(n, 4)
            else:
                rec[c] = _txt(v)
        out.append(rec)
    out.sort(key=lambda x: wbs_key(x["WBS"]) if _WBS_RE.match(x["WBS"] or "") else (10 ** 9, x["Ref"]))
    return out


def register_hash(df: pd.DataFrame) -> str:
    blob = json.dumps(canonical_register(df), sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()


def approval_record_json(df: pd.DataFrame, meta: dict) -> bytes:
    payload = {"approval": meta, "sha256": register_hash(df), "register": canonical_register(df)}
    return json.dumps(payload, indent=2, ensure_ascii=False, default=str).encode("utf-8")


def approval_record_xlsx(df: pd.DataFrame, meta: dict) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    reg = pd.DataFrame(canonical_register(df), columns=AUDIT_COLUMNS)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame([{"Field": k, "Value": str(v)} for k, v in
                      list(meta.items()) + [("Register SHA-256", register_hash(df))]]).to_excel(
            xw, sheet_name="Approval", index=False)
        reg.to_excel(xw, sheet_name="Scope register", index=False)
        head = PatternFill("solid", fgColor="0E2A47")
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head
            ws.freeze_panes = "A2"
        ws = xw.book["Scope register"]
        for col, w in zip("ABCDEFGHIJKLMN", (8, 8, 9, 60, 11, 10, 22, 11, 11, 11, 10, 60, 40, 40)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in (row[3], row[11], row[12], row[13]):
                c.alignment = Alignment(wrap_text=True, vertical="top")
        xw.book["Approval"].column_dimensions["A"].width = 24
        xw.book["Approval"].column_dimensions["B"].width = 80
    return buf.getvalue()
