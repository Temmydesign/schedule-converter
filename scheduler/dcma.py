"""DCMA 14-Point Schedule Assessment on a converted schedule (engine.Result).

Works for both the Excel converter and the document-ingestion flow. Thresholds can be set in
firm_policy.json without code changes (optional keys, defaults shown):

    "dcma_threshold_pct": 5,        checks 1, 3, 6, 8 (share of tasks / links)
    "dcma_high_float_days": 44,     check 6
    "dcma_high_duration_days": 44,  check 8
    "dcma_fs_min_pct": 90,          check 4
    "dcma_cpt_delay_days": 600      check 12 (delay injected on the critical path)

Checks 9, 10, 11 and 14 need progress, resources or a status date; at baseline they are reported
"N/A" with the reason, never as a silent pass.
"""
from __future__ import annotations

import copy
import io
from collections import Counter
from dataclasses import dataclass

import pandas as pd

from . import cpm


@dataclass
class Check:
    no: int
    name: str
    measure: str
    value: str
    threshold: str
    status: str          # Pass | Fail | N/A
    detail: str = ""


def _pct(a, b):
    return (100.0 * a / b) if b else 0.0


def dcma14(res) -> list:
    pol = res.policy or {}
    thr = float(pol.get("dcma_threshold_pct", 5))
    hf_days = int(pol.get("dcma_high_float_days", 44))
    hd_days = int(pol.get("dcma_high_duration_days", 44))
    fs_min = float(pol.get("dcma_fs_min_pct", 90))
    cpt_delay = int(pol.get("dcma_cpt_delay_days", 600))

    sched = [n for n in res.nodes if n.schedulable]
    tasks = [n for n in sched if n.kind == "task"]
    logic_pop = [n for n in sched if n.kind != "completion"]
    succ = Counter()
    for n in sched:
        for l in n.preds:
            succ[l.pred] += 1
    links = [l for n in sched for l in n.preds]

    out = []
    # 1 Logic ---------------------------------------------------------------- #
    no_pred = [n for n in logic_pop if not n.preds and n is not res.anchor]
    no_succ = [n for n in logic_pop if succ[n] == 0]
    missing = {*no_pred, *no_succ}
    p = _pct(len(missing), len(logic_pop))
    out.append(Check(1, "Logic", "Activities missing a predecessor or successor",
                     f"{len(missing)} ({p:.1f}%)", f"<= {thr:g}%", "Pass" if p <= thr else "Fail",
                     _names(missing) or "Single start (anchor) and every activity has a successor."))
    # 2 Leads ---------------------------------------------------------------- #
    leads = [l for l in links if l.lag < 0]
    out.append(Check(2, "Leads", "Relationships with negative lag", str(len(leads)), "0",
                     "Pass" if not leads else "Fail", ""))
    # 3 Lags ----------------------------------------------------------------- #
    lags = [l for l in links if l.lag > 0]
    p = _pct(len(lags), len(links))
    out.append(Check(3, "Lags", "Relationships with positive lag", f"{len(lags)} ({p:.1f}%)", f"<= {thr:g}%",
                     "Pass" if p <= thr else "Fail", ""))
    # 4 Relationship types --------------------------------------------------- #
    mix = Counter(l.type for l in links)
    fs_p = _pct(mix.get("FS", 0), len(links))
    ok4 = fs_p >= fs_min and mix.get("SF", 0) == 0
    out.append(Check(4, "Relationship types", "Finish-to-Start share (SF must be 0)",
                     f"{fs_p:.1f}% FS · " + ", ".join(f"{v} {k}" for k, v in sorted(mix.items())),
                     f">= {fs_min:g}% FS, 0 SF", "Pass" if ok4 else "Fail", ""))
    # 5 Hard constraints ----------------------------------------------------- #
    out.append(Check(5, "Hard constraints", "MSO/MFO/SNLT/FNLT constraints", "0", f"<= {thr:g}%", "Pass",
                     "Every activity is exported As Soon As Possible; dates are driven by logic only."))
    # 6 High float ----------------------------------------------------------- #
    fl = {n: res.floats.get(n) for n in tasks}
    high_f = [n for n, f in fl.items() if f is not None and f > hf_days]
    p = _pct(len(high_f), len(tasks))
    out.append(Check(6, "High float", f"Tasks with total float > {hf_days} working days",
                     f"{len(high_f)} ({p:.1f}%)", f"<= {thr:g}%", "Pass" if p <= thr else "Fail",
                     _names(high_f)))
    # 7 Negative float ------------------------------------------------------- #
    neg = [n for n, f in fl.items() if f is not None and f < 0]
    out.append(Check(7, "Negative float", "Tasks with total float < 0", str(len(neg)), "0",
                     "Pass" if not neg else "Fail", _names(neg)))
    # 8 High duration -------------------------------------------------------- #
    high_d = [n for n in tasks if n.dur > hd_days]
    p = _pct(len(high_d), len(tasks))
    out.append(Check(8, "High duration", f"Tasks longer than {hd_days} working days",
                     f"{len(high_d)} ({p:.1f}%)", f"<= {thr:g}%", "Pass" if p <= thr else "Fail",
                     _names(high_d)))
    # 9-11 need status data --------------------------------------------------- #
    out.append(Check(9, "Invalid dates", "Actuals after / forecasts before the data date", "-", "0", "N/A",
                     "Baseline schedule: no actual dates or data date yet. Assess at the first update."))
    out.append(Check(10, "Resources", "Tasks with duration but no resources/cost", "-", "<= 5%", "N/A",
                     "Resource and cost loading are outside the MVP scope."))
    out.append(Check(11, "Missed tasks", "Baseline tasks finished late", "-", "<= 5%", "N/A",
                     "Needs progress against a status date."))
    # 12 Critical path test --------------------------------------------------- #
    out.append(_cpt(res, tasks, cpt_delay))
    # 13 CPLI ---------------------------------------------------------------- #
    proj_ef = max((n.ef for n in sched if n.ef is not None), default=0)
    cpl = proj_ef
    tf_finish = 0
    cpli = (cpl + tf_finish) / cpl if cpl else 0
    out.append(Check(13, "CPLI", "(Critical path length + finish float) / critical path length",
                     f"{cpli:.2f}", ">= 0.95", "Pass" if cpli >= 0.95 else "Fail",
                     f"Critical path {cpl} working days; no finish constraint, so finish float is 0 at baseline."))
    # 14 BEI ----------------------------------------------------------------- #
    out.append(Check(14, "BEI", "Tasks completed / tasks planned to complete", "-", ">= 0.95", "N/A",
                     "Needs progress against a status date."))
    return out


def _cpt(res, tasks, delay: int) -> Check:
    crit = [n for n in tasks if (res.floats.get(n) is not None and res.floats.get(n) <= 0)]
    if not crit:
        return Check(12, "Critical path test", "Finish moves by an injected delay", "-", "Pass", "Fail",
                     "No critical task found.")
    probe = max(crit, key=lambda n: (n.dur, -n.order))
    nodes = copy.deepcopy(res.nodes)
    sched0 = [n for n in res.nodes if n.schedulable]
    f0 = max(n.ef for n in sched0)
    target = next(n for n in nodes if n.id == probe.id)
    target.dur += delay
    for n in nodes:
        n.es = n.ef = None
    try:
        cpm.forward_pass(nodes, 0)
    except cpm.CycleError as exc:
        return Check(12, "Critical path test", "Finish moves by an injected delay", "-", "Pass", "Fail", str(exc))
    f1 = max(n.ef for n in nodes if n.schedulable)
    moved = f1 - f0
    ok = moved == delay
    return Check(12, "Critical path test", f"Delay of {delay} working days added to '{probe.name[:50]}'",
                 f"finish moved {moved} days", f"= {delay}", "Pass" if ok else "Fail",
                 "The critical path is continuous from start to finish." if ok else
                 "The critical path is broken (constraint, open end or loose logic).")


def _names(nodes, limit: int = 6) -> str:
    nodes = sorted(nodes, key=lambda n: n.order)
    txt = "; ".join(f"{n.id} {n.name[:40]}" for n in nodes[:limit])
    return txt + (f" … (+{len(nodes) - limit})" if len(nodes) > limit else "")


def dcma_frame(checks: list) -> pd.DataFrame:
    return pd.DataFrame([{"#": c.no, "Check": c.name, "Measure": c.measure, "Result": c.value,
                          "Threshold": c.threshold, "Status": c.status, "Detail": c.detail} for c in checks])


def summary(checks: list) -> dict:
    s = Counter(c.status for c in checks)
    return {"pass": s.get("Pass", 0), "fail": s.get("Fail", 0), "na": s.get("N/A", 0),
            "assessed": s.get("Pass", 0) + s.get("Fail", 0)}


def dcma_xlsx(res, checks: list) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    buf = io.BytesIO()
    df = dcma_frame(checks)
    sm = summary(checks)
    head = pd.DataFrame([
        {"Item": "Project", "Value": res.project.name},
        {"Item": "Calendar", "Value": res.calendar_name},
        {"Item": "Start / finish", "Value": f"{res.stats['output_start']:%d %b %Y} / {res.stats['output_finish']:%d %b %Y}"},
        {"Item": "Checks passed", "Value": f"{sm['pass']} of {sm['assessed']} assessed ({sm['na']} N/A at baseline)"},
    ])
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        head.to_excel(xw, sheet_name="DCMA-14", index=False, startrow=0)
        df.to_excel(xw, sheet_name="DCMA-14", index=False, startrow=len(head) + 2)
        ws = xw.book["DCMA-14"]
        fill = PatternFill("solid", fgColor="0E2A47")
        for row_idx in (1, len(head) + 3):
            for c in ws[row_idx]:
                if c.value is not None:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = fill
        colors = {"Pass": "D9F2E3", "Fail": "FBE0DE", "N/A": "EEF1F5"}
        for r in range(len(head) + 4, len(head) + 4 + len(df)):
            st = ws.cell(r, 6).value
            ws.cell(r, 6).fill = PatternFill("solid", fgColor=colors.get(st, "FFFFFF"))
            ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="top")
        for col, w in zip("ABCDEFG", (16, 22, 48, 26, 16, 9, 70)):
            ws.column_dimensions[col].width = w
    return buf.getvalue()
