"""Regression tests on synthetic schedules in many layouts (no client data in the repo).

Run:  pytest -q
"""
import datetime as dt
import io
import os
import sys
import xml.etree.ElementTree as ET

import numpy as np
import openpyxl
import pytest
from openpyxl.styles import Alignment

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from scheduler import exporters  # noqa: E402
from scheduler.engine import Options, convert  # noqa: E402

D0 = dt.date(2026, 11, 2)  # a Monday


def wd(n):
    """Date n working days after D0 (n=0 -> D0)."""
    return np.busday_offset(np.datetime64(D0), n, roll="forward").astype(dt.date)


def book(rows, title=None, sheet="Schedule", indents=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    r0 = 1
    if title:
        ws.cell(1, 1, title)
        r0 = 3
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ws.cell(r0 + i, j + 1, v)
            if indents and i > 0 and j == 0 and indents[i - 1]:
                c.alignment = Alignment(indent=indents[i - 1])
    ws.cell(r0 + len(rows) + 2, 2, "Legend: blue = task")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def check(res):
    st = res.stats
    assert st["date_mismatches"] == 0, res.issues
    assert st["finish_matches"]
    assert st["open_starts"] == 0
    assert res.nodes[0].kind == "project" and res.nodes[0].level == 1
    for n in res.nodes:
        if n.kind == "summary":
            assert n.children[-1].kind == "completion", n.name
            assert n.children[-1].name.endswith("completed")
        if n.schedulable and n is not res.anchor:
            assert n.preds, n.name
        if n.kind != "project":
            assert n.level == n.parent.level + 1
    ET.fromstring(exporters.to_mspdi(res))                 # well-formed XML
    csv = exporters.to_csv(res).decode("utf-8-sig")
    assert csv.splitlines()[0] == "ID,Task Name,Duration,Start Date,Finish Date,Predecessors,Outline Level,WBS"
    assert exporters.to_qa_xlsx(res)[:2] == b"PK"


def test_system_grouped_weeks():
    rows = [["System", "Activity", "Start Wk", "Dur", "Finish Wk", "Start date", "Finish date"],
            ["PRG", "Mobilisation", 1, 5, 5, wd(0), wd(24)],
            ["PRG", "Project management", 1, 20, 20, wd(0), wd(99)],
            ["A", "A - Data review", 1, 4, 4, wd(0), wd(19)],
            ["A", "A - Design", 3, 6, 8, wd(10), wd(39)],
            ["A", "A - Report", 9, 2, 10, wd(40), wd(49)],
            ["B", "B - Data review", 1, 6, 6, wd(0), wd(29)],
            ["B", "B - Design", 5, 8, 12, wd(20), wd(59)],
            ["MS", "DG1 - Design approval (all systems)", 12, 0, 12, wd(55), wd(59)]]
    res = convert(book(rows, "Programme | Test Study - Summary | Prepared by X", "Integrated L2 Schedule"),
                  "t.xlsx", Options())
    check(res)
    assert res.info["mode"] == "group"
    names = [n.name for n in res.nodes]
    assert "System: PRG" in names and "PRG completed" in names and "Programme Milestones completed" in names
    assert res.nodes[2].name == "Mobilisation" and not res.nodes[2].preds      # anchor
    dur = {n.name: exporters.duration_text(n) for n in res.nodes}
    assert dur["A - Design"] == "6 wks"


def test_wbs_type_columns():
    rows = [["WBS", "Activity", "Type", "Dur (wks)", "Start date", "Finish date"],
            ["1.0", "MOBILISATION", "Sum", 4, wd(0), wd(19)],
            ["1.1", "Notice to Proceed", "Mstone", 0, wd(0), wd(4)],
            ["1.2", "Mobilise team", "Task", 2, wd(0), wd(9)],
            ["1.3", "PEP", "Task", 2, wd(10), wd(19)],
            ["2.0", "ENGINEERING", "Sum", 8, wd(10), wd(49)],
            ["2.1", "Design basis", "Task", 3, wd(10), wd(24)],
            ["2.1.1", "Design basis review", "Review", 1, wd(25), wd(29)],
            ["2.2", "Detailed design", "Task", 4, wd(30), wd(49)],
            ["2.3", "DG2 - Design approval", "Mstone", 0, wd(45), wd(49)]]
    res = convert(book(rows), "w.xlsx", Options(project_name="Test Project"))
    check(res)
    assert res.info["mode"] == "wbs"
    assert res.calendar_name == "Test Project CALENDAR"
    lv = {n.name: n.level for n in res.nodes}
    assert lv["MOBILISATION"] == 2 and lv["Design basis"] == 3 and lv["Design basis review"] == 4
    assert "Design basis completed" in lv   # nested summary also gets its milestone


def test_flat_days_text_dates_dayfirst():
    rows = [["Task Name", "Duration (days)", "Start", "End"],
            ["Survey", 5, "02/11/2026", "06/11/2026"],
            ["Design", 10, "09/11/2026", "20/11/2026"],
            ["Procure", 15, "16/11/2026", "04/12/2026"],
            ["Install", 10, "07/12/2026", "18/12/2026"],
            ["Handover", 0, "18/12/2026", "18/12/2026"]]
    res = convert(book(rows), "f.xlsx", Options(dayfirst=True))
    check(res)
    assert res.stats["output_finish"] == dt.date(2026, 12, 18)
    assert exporters.duration_text(next(n for n in res.nodes if n.name == "Design")) == "10 days"


def test_source_predecessors_are_used():
    rows = [["ID", "Task", "Duration", "Start", "Finish", "Predecessors"],
            [1, "Start", 0, D0, D0, None],
            [2, "Dig", "5d", wd(0), wd(4), "1"],
            [3, "Pour", "5d", wd(7), wd(11), "2FS+2d"],
            [4, "Cure", "3d", wd(12), wd(14), "3"]]
    res = convert(book(rows), "p.xlsx", Options())
    check(res)
    pour = next(n for n in res.nodes if n.name == "Pour")
    assert (pour.preds[0].pred.name, pour.preds[0].type, pour.preds[0].lag) == ("Dig", "FS", 2)
    assert "spreadsheet" in res.info["logic_source"]


def test_gantt_only_matrix():
    weeks = [wd(5 * i) for i in range(8)]
    header = ["Activity"] + list(range(1, 9))
    rows = [[None] + weeks, header,
            ["Kick-off", "◆"] + [None] * 7,
            ["Studies", 1, 1, 1] + [None] * 5,
            ["Design", None, None, 1, 1, 1, None, None, None],
            ["Review", None, None, None, None, None, 1, 1, None],
            ["Complete", None, None, None, None, None, None, None, "◆"]]
    res = convert(book(rows), "g.xlsx", Options())
    check(res)
    assert next(n for n in res.nodes if n.name == "Kick-off").kind == "milestone"


def test_indent_hierarchy_and_heading_rows():
    rows = [["Activity", "Start", "Finish"],
            ["Phase 1", None, None],
            ["Task A", wd(0), wd(9)],
            ["Task B", wd(10), wd(19)],
            ["Phase 2", None, None],
            ["Task C", wd(20), wd(29)]]
    res = convert(book(rows, indents=[0, 1, 1, 0, 1]), "i.xlsx", Options())
    check(res)
    assert res.info["mode"] == "indent"
    assert [n.name for n in res.nodes if n.kind == "completion"] == ["Phase 1 completed", "Phase 2 completed"]


def test_csv_input_and_rebaseline():
    text = "Activity,Start date,Finish date\nA,2026-11-02,2026-11-13\nB,2026-11-16,2026-11-27\n"
    res = convert(text.encode(), "c.csv", Options(project_start=dt.date(2027, 1, 4)))
    check(res)
    assert res.stats["output_start"] == dt.date(2027, 1, 4)
    assert res.stats["output_finish"] == dt.date(2027, 1, 29)


def test_holidays_respected():
    rows = [["Activity", "Start", "Finish"],
            ["A", dt.date(2026, 12, 21), dt.date(2026, 12, 31)],   # spans Christmas / Boxing Day
            ["B", dt.date(2027, 1, 4), dt.date(2027, 1, 8)]]
    hol = [dt.date(2026, 12, 25), dt.date(2026, 12, 28)]
    res = convert(book(rows), "h.xlsx", Options(holidays=hol))
    check(res)
    a = next(n for n in res.nodes if n.name == "A")
    assert a.dur == 7
    assert b"<DayType>0</DayType>" in exporters.to_mspdi(res)


def test_no_schedule_raises():
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Phone"])
    wb.active.append(["Ada", "123"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError):
        convert(buf.getvalue(), "x.xlsx", Options())


def test_planner_edit_is_reverified():
    from scheduler.engine import apply_edited_predecessors
    rows = [["Activity", "Start", "Finish"], ["A", wd(0), wd(9)], ["B", wd(10), wd(19)]]
    res = convert(book(rows), "e.xlsx", Options())
    b = next(n for n in res.nodes if n.name == "B")
    res = apply_edited_predecessors(res, {b.id: f"{b.id - 1}FS+1 wk"})
    assert res.stats["date_mismatches"] == 1 and not res.stats["finish_matches"]


def test_close_open_ends_keeps_dates():
    rows = [["Activity", "Start", "Finish"], ["A", wd(0), wd(9)], ["B", wd(0), wd(4)], ["C", wd(10), wd(19)]]
    res = convert(book(rows), "o.xlsx", Options(policy={"close_open_ends": True,
                                                        "completion_milestone_for_project_title": True}))
    check(res)
    assert res.stats["open_ends"] == 0
